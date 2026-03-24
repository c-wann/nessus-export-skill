#!/usr/bin/env python3
"""
Nessus Export Tool — bundled skill script
-----------------------------------------
Logs into the Nessus web portal via Chrome browser automation,
exports a scan as .nessus XML, scrapes the vulnerability table,
and produces a colour-coded Excel + clean .nessus file.

Usage:
  python3 nessus_export_tool.py --scan 19 --host 2 --label GETS_PRD_168
  python3 nessus_export_tool.py --scan 19 --host 2 --user admin --pass admin
"""

import argparse, getpass, os, re, shutil, sys, time
import xml.etree.ElementTree as ET
from collections import defaultdict, Counter
from datetime import datetime
from pathlib import Path

import urllib3
urllib3.disable_warnings()

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException
from webdriver_manager.chrome import ChromeDriverManager

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Constants ──────────────────────────────────────────────────────────────────
SEV_MAP   = {"0": "Info", "1": "Low", "2": "Medium", "3": "High", "4": "Critical"}
SEV_COLOR = {
    "Critical": "FF0000", "High": "FF6600", "Medium": "FFCC00",
    "Low": "66CC00", "Info": "99CCFF", "Mixed": "CC99FF",
}

# ── Helpers ────────────────────────────────────────────────────────────────────
def sfill(sev): return PatternFill("solid", fgColor=SEV_COLOR.get(sev, "FFFFFF"))

thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"))

def hdr_row(ws, row_num, cols):
    hf = Font(bold=True, color="FFFFFF", size=11)
    hb = PatternFill("solid", fgColor="2F4F8F")
    for ci, h in enumerate(cols, 1):
        c = ws.cell(row_num, ci, h)
        c.font = hf; c.fill = hb
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = thin_border
    ws.row_dimensions[row_num].height = 30

def wcell(ws, row, ci, val, fill):
    c = ws.cell(row, ci, val)
    c.fill = fill; c.border = thin_border
    c.alignment = Alignment(wrap_text=True, vertical="top")

def set_col_widths(ws, widths: dict):
    for ci, w in widths.items():
        ws.column_dimensions[get_column_letter(ci)].width = w

# ── Browser ────────────────────────────────────────────────────────────────────
def make_driver(download_dir: str) -> webdriver.Chrome:
    opts = Options()
    opts.add_argument("--ignore-certificate-errors")
    opts.add_argument("--ignore-ssl-errors")
    opts.add_argument("--headless=new")
    opts.add_argument("--window-size=1920,1080")
    opts.add_argument("--no-sandbox")
    opts.add_argument("--disable-dev-shm-usage")
    prefs = {
        "download.default_directory": download_dir,
        "download.prompt_for_download": False,
        "download.directory_upgrade": True,
        "safebrowsing.enabled": False,
    }
    opts.add_experimental_option("prefs", prefs)
    opts.add_experimental_option("excludeSwitches", ["enable-logging"])
    return webdriver.Chrome(
        service=Service(ChromeDriverManager().install()), options=opts)


def login(driver, base_url: str, user: str, pwd: str):
    print("[*] Opening Nessus portal…")
    driver.get(base_url); time.sleep(3)
    print("[*] Logging in…")
    driver.find_element(
        By.XPATH, "//input[@type='text' or @name='username' or @id='username']"
    ).send_keys(user)
    driver.find_element(By.XPATH, "//input[@type='password']").send_keys(pwd)
    driver.find_element(By.XPATH,
        "//button[@type='submit' or contains(.,'Sign In') or contains(.,'Log In')]"
    ).click()
    time.sleep(4)
    if "login" in driver.current_url.lower():
        raise RuntimeError("Login failed — check username/password")
    print("[+] Logged in")


def export_nessus_xml(driver, base_url: str, scan_id: int,
                      download_dir: str) -> str:
    """Click Export → Nessus in the UI, wait for download, return path."""
    print(f"[*] Navigating to scan {scan_id} report…")
    driver.get(f"{base_url}/#/scans/reports/{scan_id}/hosts"); time.sleep(4)

    # Dismiss notification overlays
    driver.execute_script(
        "document.querySelectorAll('.notification-message,.notification,.toast')"
        ".forEach(el=>el.remove());")
    time.sleep(1)

    export_span = WebDriverWait(driver, 15).until(
        EC.presence_of_element_located((By.ID, "export")))
    driver.execute_script("arguments[0].click();", export_span)
    time.sleep(1)

    nessus_li = WebDriverWait(driver, 10).until(
        EC.presence_of_element_located(
            (By.XPATH, "//span[@id='export']//li[@data-value='nessus']")))
    driver.execute_script("arguments[0].click();", nessus_li)
    print("[*] Waiting for .nessus download…")

    path = _wait_for_download(download_dir, ".nessus", 60)
    if not path:
        raise RuntimeError("Download timed out after 60s")
    print(f"[+] Downloaded: {Path(path).name}")
    return path


def scrape_ui_table(driver, base_url: str, scan_id: int,
                    host_id: int) -> list[dict]:
    """Scrape the vulnerability table exactly as shown in the Nessus UI."""
    url = f"{base_url}/#/scans/reports/{scan_id}/hosts/{host_id}/vulnerabilities"
    print(f"[*] Scraping vulnerability table: {url}")
    driver.get(url); time.sleep(5)

    rows = driver.find_elements(By.XPATH, "//table//tbody//tr")
    print(f"[+] UI shows {len(rows)} vulnerability rows")

    ui_rows = []
    for row in rows:
        cells  = row.find_elements(By.TAG_NAME, "td")
        texts  = [c.text.strip() for c in cells]
        cls    = row.get_attribute("class") or ""
        pid    = row.get_attribute("data-id") or ""
        sev_n  = row.get_attribute("data-severity") or "0"
        is_grp = "group" in cls

        name_raw = texts[5] if len(texts) > 5 else ""
        if "\n" in name_raw:
            name_raw = name_raw.split("\n", 1)[1]

        raw_sev = texts[1].strip() if len(texts) > 1 else ""
        if raw_sev.upper() == "MIXED":
            sev_label = "Mixed"
        elif raw_sev.upper() in ("CRITICAL","HIGH","MEDIUM","LOW","INFO"):
            sev_label = raw_sev.title()
        else:
            sev_label = SEV_MAP.get(sev_n, "Info")

        ui_rows.append({
            "plugin_id":    pid,
            "name":         name_raw,
            "severity":     sev_label,
            "severity_num": int(sev_n),
            "cvss":         texts[2] if len(texts) > 2 else "",
            "vpr":          texts[3] if len(texts) > 3 else "",
            "epss":         texts[4] if len(texts) > 4 else "",
            "family":       texts[6] if len(texts) > 6 else "",
            "count":        texts[7] if len(texts) > 7 else "1",
            "is_group":     is_grp,
        })
    return ui_rows


def _wait_for_download(directory: str, ext: str, timeout: int = 60):
    deadline = time.time() + timeout
    while time.time() < deadline:
        files = sorted(
            [f for f in Path(directory).glob(f"*{ext}")
             if not f.name.endswith(".crdownload")],
            key=lambda f: f.stat().st_mtime, reverse=True)
        if files and time.time() - files[0].stat().st_mtime < 120:
            return str(files[0])
        time.sleep(1)
    return None

# ── XML Parser ─────────────────────────────────────────────────────────────────
def parse_nessus(path: str) -> tuple[dict, dict]:
    """Parse .nessus XML → (by_plugin_id dict, host_props dict)."""
    print(f"[*] Parsing {Path(path).name}…")
    tree = ET.parse(path)
    root = tree.getroot()

    by_plugin  = defaultdict(list)
    host_props = {}

    for rh in root.iter("ReportHost"):
        props = {t.get("name"): t.text for t in rh.findall("HostProperties/tag")}
        if not host_props:
            host_props = props
        for item in rh.findall("ReportItem"):
            pid = item.get("pluginID", "")
            sev = item.get("severity", "0")
            cve = ", ".join(c.text for c in item.findall("cve") if c.text)
            bid = ", ".join(c.text for c in item.findall("bid") if c.text)
            by_plugin[pid].append({
                "plugin_id":     pid,
                "plugin_name":   item.get("pluginName", ""),
                "severity":      SEV_MAP.get(sev, "Info"),
                "severity_num":  int(sev),
                "port":          item.get("port", ""),
                "protocol":      item.get("protocol", ""),
                "svc_name":      item.get("svc_name", ""),
                "family":        item.get("pluginFamily", ""),
                "cvss":          (item.findtext("cvss_base_score") or
                                  item.findtext("cvss3_base_score") or ""),
                "cvss_vector":   (item.findtext("cvss_base_vector") or
                                  item.findtext("cvss3_vector") or ""),
                "cve":           cve,
                "bid":           bid,
                "synopsis":      (item.findtext("synopsis") or "").strip(),
                "description":   (item.findtext("description") or "").strip(),
                "solution":      (item.findtext("solution") or "").strip(),
                "plugin_output": (item.findtext("plugin_output") or "").strip(),
                "see_also":      (item.findtext("see_also") or "").strip(),
                "patch_date":    (item.findtext("patch_publication_date") or ""),
            })

    total = sum(len(v) for v in by_plugin.values())
    print(f"[+] Parsed {total} findings across {len(by_plugin)} unique plugins")
    return by_plugin, host_props

# ── Excel Builder ──────────────────────────────────────────────────────────────
def build_excel(ui_rows, by_plugin, host_props,
                scan_id, host_id, label, out_path):
    print("[*] Building Excel workbook…")
    wb = openpyxl.Workbook()

    ip   = host_props.get("host-ip", "N/A")
    fqdn = host_props.get("host-fqdn", ip)
    mac  = host_props.get("mac-address", host_props.get("mac_address", "N/A"))
    os_  = host_props.get("operating-system", host_props.get("os", "N/A"))
    all_findings = sorted(
        [f for items in by_plugin.values() for f in items],
        key=lambda f: (-f["severity_num"], f["plugin_name"]))

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws0 = wb.active; ws0.title = "Summary"
    title_font = Font(bold=True, size=14)
    ws0["A1"] = "Nessus Vulnerability Report"; ws0["A1"].font = title_font
    ws0.merge_cells("A1:C1")

    info = [
        ("Generated",     datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Scan ID",       str(scan_id)),
        ("Host ID",       str(host_id)),
        ("Host Label",    label),
        ("IP Address",    ip),
        ("FQDN",          fqdn),
        ("MAC Address",   mac),
        ("OS",            os_),
        ("UI Vuln Groups",str(len(ui_rows))),
        ("Total Findings",str(len(all_findings))),
    ]
    for ri, (k, v) in enumerate(info, 2):
        ws0.cell(ri, 1, k).font = Font(bold=True)
        ws0.cell(ri, 2, v)

    ws0["A13"] = "Severity"; ws0["B13"] = "Groups"; ws0["C13"] = "Findings"
    for c in [ws0["A13"], ws0["B13"], ws0["C13"]]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F4F8F")
        c.alignment = Alignment(horizontal="center")

    grp_cnt = Counter(r["severity"] for r in ui_rows)
    fnd_cnt = Counter(f["severity"] for f in all_findings)
    for ri2, sev in enumerate(["Critical","High","Medium","Low","Info","Mixed"], 14):
        if grp_cnt.get(sev, 0) or fnd_cnt.get(sev, 0):
            ws0.cell(ri2, 1, sev).fill = sfill(sev)
            ws0.cell(ri2, 1).font = Font(bold=True)
            ws0.cell(ri2, 2, grp_cnt.get(sev, 0))
            ws0.cell(ri2, 3, fnd_cnt.get(sev, 0))

    ws0.column_dimensions["A"].width = 20
    ws0.column_dimensions["B"].width = 12
    ws0.column_dimensions["C"].width = 14

    # ── Sheet 2: Vulnerabilities (UI groups) ──────────────────────────────────
    n_ui   = len(ui_rows)
    n_fnd  = len(all_findings)
    ws1 = wb.create_sheet(f"Vulnerabilities ({n_ui})")
    cols1 = ["#","Severity","CVSS","VPR","EPSS","Vulnerability Name",
             "Family","Count","Type","Plugin ID","CVE",
             "Synopsis","Solution","Description","Plugin Output",
             "Port/Protocol","Patch Date","See Also"]
    hdr_row(ws1, 1, cols1); ws1.freeze_panes = "A2"

    for ri, row in enumerate(ui_rows, 2):
        pid    = row["plugin_id"]
        items  = by_plugin.get(pid, [])
        d      = items[0] if items else {}
        ports  = ", ".join(sorted(set(
            f"{i['port']}/{i['protocol']}" for i in items if i.get("port"))))
        p_out  = "\n---\n".join(
            i["plugin_output"] for i in items if i.get("plugin_output"))
        cvss   = row["cvss"] if row["cvss"] not in ("...","") else d.get("cvss","")
        fill   = sfill(row["severity"])
        for ci, v in enumerate([
            ri-1, row["severity"], cvss, row["vpr"], row["epss"],
            row["name"], row["family"], row["count"],
            "Group" if row["is_group"] else "Single", pid,
            d.get("cve",""), d.get("synopsis",""), d.get("solution",""),
            d.get("description",""), p_out, ports,
            d.get("patch_date",""), d.get("see_also",""),
        ], 1):
            wcell(ws1, ri, ci, v, fill)

    set_col_widths(ws1, {1:4,2:10,3:8,4:7,5:8,6:45,7:28,8:7,9:7,
                         10:9,11:22,12:45,13:45,14:60,15:60,16:18,17:12,18:40})

    # ── Sheet 3: All Findings ─────────────────────────────────────────────────
    ws2 = wb.create_sheet(f"All Findings ({n_fnd})")
    cols2 = ["#","Severity","CVSS","Plugin ID","Plugin Name","Family",
             "Port","Protocol","Service","CVE","Synopsis",
             "Solution","Description","Plugin Output","Patch Date"]
    hdr_row(ws2, 1, cols2); ws2.freeze_panes = "A2"

    for ri, f in enumerate(all_findings, 2):
        fill = sfill(f["severity"])
        for ci, v in enumerate([
            ri-1, f["severity"], f["cvss"], f["plugin_id"], f["plugin_name"],
            f["family"], f["port"], f["protocol"], f["svc_name"],
            f["cve"], f["synopsis"], f["solution"], f["description"],
            f["plugin_output"], f["patch_date"],
        ], 1):
            wcell(ws2, ri, ci, v, fill)

    set_col_widths(ws2, {1:4,2:10,3:8,4:9,5:45,6:28,7:7,8:10,9:12,
                         10:22,11:45,12:45,13:60,14:60,15:12})

    # ── Sheet 4: Critical & High ──────────────────────────────────────────────
    ws3 = wb.create_sheet("Critical & High")
    hdr_row(ws3, 1, cols2); ws3.freeze_panes = "A2"
    ch = [f for f in all_findings if f["severity"] in ("Critical","High")]
    for ri, f in enumerate(ch, 2):
        fill = sfill(f["severity"])
        for ci, v in enumerate([
            ri-1, f["severity"], f["cvss"], f["plugin_id"], f["plugin_name"],
            f["family"], f["port"], f["protocol"], f["svc_name"],
            f["cve"], f["synopsis"], f["solution"], f["description"],
            f["plugin_output"], f["patch_date"],
        ], 1):
            wcell(ws3, ri, ci, v, fill)
    set_col_widths(ws3, {1:4,2:10,3:8,4:9,5:45,6:28,7:7,8:10,9:12,
                         10:22,11:45,12:45,13:60,14:60,15:12})

    wb.save(out_path)
    print(f"[+] Saved Excel → {out_path}")
    return out_path

# ── Main ───────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(
        description="Export Nessus vulnerabilities to Excel + .nessus via browser UI")
    parser.add_argument("--url",    default="https://127.0.0.1:8834",
                        help="Nessus portal URL (default: https://127.0.0.1:8834)")
    parser.add_argument("--user",   default="admin", help="Username (default: admin)")
    parser.add_argument("--pass",   dest="password",  help="Password")
    parser.add_argument("--scan",   type=int, required=True, help="Scan ID (from URL)")
    parser.add_argument("--host",   type=int, required=True, help="Host ID (from URL)")
    parser.add_argument("--label",  default="",
                        help="Label for output filenames (e.g. GETS_PRD_168)")
    parser.add_argument("--outdir", default=str(Path.home()),
                        help="Output directory (default: home dir)")
    args = parser.parse_args()

    pwd   = args.password or getpass.getpass(f"Password for '{args.user}': ")
    label = args.label or f"scan{args.scan}_host{args.host}"
    # Sanitise label for filenames
    label = re.sub(r"[^\w\-]", "_", label)

    out_dir    = str(Path(args.outdir).expanduser())
    xlsx_path  = os.path.join(out_dir, f"nessus_{label}_vulnerabilities.xlsx")
    nessus_path= os.path.join(out_dir, f"nessus_{label}_vulnerabilities.nessus")

    driver = make_driver(out_dir)
    try:
        login(driver, args.url, args.user, pwd)
        tmp_nessus = export_nessus_xml(driver, args.url, args.scan, out_dir)
        ui_rows    = scrape_ui_table(driver, args.url, args.scan, args.host)
    except RuntimeError as e:
        print(f"[✗] {e}"); sys.exit(1)
    finally:
        driver.quit()

    # Save clean .nessus file
    shutil.copy2(tmp_nessus, nessus_path)
    print(f"[+] Saved .nessus → {nessus_path}")

    by_plugin, host_props = parse_nessus(nessus_path)
    if not by_plugin:
        print("[!] No findings in XML — check scan/host IDs"); sys.exit(1)

    build_excel(ui_rows, by_plugin, host_props,
                args.scan, args.host, label, xlsx_path)

    # Summary
    grp_cnt = Counter(r["severity"] for r in ui_rows)
    fnd_cnt = Counter(f["severity"]
                      for items in by_plugin.values() for f in items)

    print(f"\n{'═'*48}")
    print(f"  {label}  —  Scan {args.scan} / Host {args.host}")
    print(f"{'═'*48}")
    print(f"  {'Severity':<12} {'UI Groups':>10}  {'Findings':>10}")
    print(f"  {'-'*36}")
    for sev in ["Critical","High","Medium","Low","Info","Mixed"]:
        g, f = grp_cnt.get(sev,0), fnd_cnt.get(sev,0)
        if g or f:
            print(f"  {sev:<12} {g:>10}  {f:>10}")
    print(f"  {'-'*36}")
    print(f"  {'TOTAL':<12} {len(ui_rows):>10}  {sum(fnd_cnt.values()):>10}")
    print(f"{'═'*48}")
    print(f"\n[✓] Excel  → {xlsx_path}")
    print(f"[✓] Nessus → {nessus_path}")
    print(f"\n    open '{xlsx_path}'")


if __name__ == "__main__":
    main()
